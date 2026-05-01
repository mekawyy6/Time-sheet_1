from __future__ import annotations

import calendar
import json
from collections import OrderedDict, defaultdict
from datetime import date, datetime, time
from io import BytesIO
from typing import Annotated

from fastapi.responses import Response, RedirectResponse
from fastapi import APIRouter, Depends, Form, Request, status
from fastapi.responses import RedirectResponse, Response, StreamingResponse
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.core.constants import (
    ACTION_OPTIONS,
    AREAS,
    RESERVATION_TYPES,
    SITE_TYPES,
    STAY_PLACES,
    TOWER_TYPES,
    TRAVEL_TYPES,
    YES_NO,
)
from app.deps import get_db
from app.models.daily_entry import DailyEntry
from app.models.site_evaluation import SiteEvaluation
from app.models.user import User

router = APIRouter()
templates = Jinja2Templates(directory="app/templates")


# ---------- General helpers ----------
def current_month_label() -> str:
    return datetime.now().strftime("%B %Y")


def current_month_value() -> str:
    return datetime.now().strftime("%Y-%m")


def get_profile(request: Request, db: Session) -> User | None:
    profile_id = request.cookies.get("profile_id")

    if profile_id and profile_id.isdigit():
        profile = db.query(User).filter(User.id == int(profile_id)).first()
        if profile:
            return profile

    return None


def profile_ready(profile: User | None) -> bool:
    return bool(profile and profile.name and profile.employee_id)


def ensure_profile(request: Request, db: Session):
    profile = get_profile(request, db)

    if not profile:
        return RedirectResponse(url="/setup", status_code=302)

    return profile


def parse_time_value(value: str | None) -> time | None:
    if not value:
        return None
    try:
        return time.fromisoformat(value)
    except ValueError:
        return None


def time_to_input(value: time | None) -> str:
    return value.strftime("%H:%M") if value else ""


def time_to_sheet(value: time | None) -> str:
    if not value:
        return ""
    hour = value.hour % 12 or 12
    suffix = "AM" if value.hour < 12 else "PM"
    return f"{hour}:{value.minute:02d}{suffix}"


def to_float(value: str | None) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return float(value)
    except ValueError:
        return 0.0


def calculate_duration_text(start: time, end: time) -> str:
    start_minutes = start.hour * 60 + start.minute
    end_minutes = end.hour * 60 + end.minute
    diff = max(end_minutes - start_minutes, 0)
    hours, minutes = divmod(diff, 60)
    if minutes == 0:
        return str(hours)
    return f"{hours}:{minutes:02d}"


def calc_eval_scores(payload: dict[str, str], site_type: str | None) -> dict[str, float]:
    rru_qty = to_float(payload.get("rru_qty"))
    rru_factor = 1.5 if site_type == "Huawei" else 2.0
    rru_score = rru_qty * rru_factor
    fiber_score = to_float(payload.get("fiber_qty")) * 0.5
    dc_score = to_float(payload.get("dc_qty")) * 0.6
    indoor_score = to_float(payload.get("indoor_hu_survey_psu_qty")) * 0.25
    sw_er_score = to_float(payload.get("sw_er_qty")) * 0.75
    power_cabinet_score = 5.0 if payload.get("power_cabinet") == "Yes" else 0.0
    batteries_addition_score = 3.0 if payload.get("batteries_addition") == "Yes" else 0.0
    subtotal = rru_score + fiber_score + dc_score + indoor_score + sw_er_score + power_cabinet_score + batteries_addition_score
    team_members_qty = max(int(to_float(payload.get("team_members_qty"))), 1)
    final_score = subtotal / team_members_qty
    return {
        "rru_score": rru_score,
        "fiber_score": fiber_score,
        "dc_score": dc_score,
        "indoor_hu_survey_psu_score": indoor_score,
        "sw_er_score": sw_er_score,
        "power_cabinet_score": power_cabinet_score,
        "batteries_addition_score": batteries_addition_score,
        "adjusted_total": subtotal,
        "final_score": final_score,
    }


def duplicate_entry_exists(db: Session, user_id: int, entry_date: date, exclude_id: int | None = None):
    query = db.query(DailyEntry).filter(
        DailyEntry.user_id == user_id,
        DailyEntry.entry_date == entry_date
    )

    if exclude_id is not None:
        query = query.filter(DailyEntry.id != exclude_id)

    return db.query(query.exists()).scalar()


def month_bounds(month_value: str):
    if month_value and len(month_value) == 7:
        y, m = month_value.split('-')
        year, month = int(y), int(m)
    else:
        now = datetime.now()
        year, month = now.year, now.month
    start = date(year, month, 1)
    end = date(year + (month == 12), 1 if month == 12 else month + 1, 1)
    return start, end


def month_days(month_value: str):
    start, _ = month_bounds(month_value)
    _, last_day = calendar.monthrange(start.year, start.month)
    return [date(start.year, start.month, d) for d in range(1, last_day + 1)]


def sheet_date(d: date) -> str:
    return d.strftime("%d-%b-%y").lstrip("0")


def sheet_day(d: date) -> str:
    return d.strftime("%A")


def map_travel_flags(value: str | None) -> tuple[str, str]:
    value = (value or "").strip().lower()
    if value in {"stay"}:
        return "1", ""
    if value in {"back to cairo", "travel without stay"}:
        return "", "1"
    return "", ""


def fill_headers(ws, title: str, profile: User, month_value: str, total_columns: int):
    month_start, _ = month_bounds(month_value)
    ws.merge_cells(start_row=1, start_column=11, end_row=1, end_column=total_columns)
    title_cell = ws.cell(1, 11)
    title_cell.value = f"{month_start.month}/{month_start.day} /{month_start.year}  {title}"
    title_cell.font = Font(bold=True, size=14)
    title_cell.fill = PatternFill("solid", fgColor="FFF200")
    title_cell.alignment = Alignment(horizontal="center")

    ws.cell(3, 1, "Name")
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=8)
    ws.cell(3, 2, profile.name)
    ws.cell(4, 1, "ID")
    ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=5)
    ws.cell(4, 2, profile.employee_id)
    for row in (3, 4):
        for col in range(1, total_columns + 1):
            ws.cell(row, col).font = Font(bold=True)


def style_sheet(ws, header_row: int, total_columns: int):
    thin = Side(style="thin", color="000000")
    red_fill = PatternFill("solid", fgColor="FF0000")
    grey_fill = PatternFill("solid", fgColor="D9D9D9")
    for col in range(1, total_columns + 1):
        cell = ws.cell(header_row, col)
        cell.font = Font(bold=True, color="000000")
        cell.fill = red_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=header_row, max_row=ws.max_row, min_col=1, max_col=total_columns):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if cell.row > header_row and cell.row % 2 == 1:
                cell.fill = grey_fill
    for col in range(1, total_columns + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16


def build_monthly_export(entries: list[DailyEntry], eval_map: dict[int, SiteEvaluation], profile: User, month_value: str) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    entry_by_date = {e.entry_date: e for e in entries}
    days = month_days(month_value)

    # Styles
    header_font = Font(bold=True, color="000000")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    header_fill = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    def apply_sheet_format(ws, header_row, total_cols):
        # Body formatting
        for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row, min_col=1, max_col=total_cols):
            for cell in row:
                try:
                    cell.alignment = center_align
                    cell.border = thin_border
                except Exception:
                    pass

        # Left align some long-text columns if they exist
        for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
            for col_idx in [8, 9, 17, 23]:
                if col_idx <= total_cols:
                    try:
                        row[col_idx - 1].alignment = left_align
                    except Exception:
                        pass

        # Freeze under table header
        ws.freeze_panes = f"A{header_row + 1}"

        # Auto width بدون المرور على merged cells
        for col_idx in range(1, total_cols + 1):
            col_letter = ws.cell(header_row, col_idx).column_letter
            max_length = 0
            for row_idx in range(1, ws.max_row + 1):
                cell = ws.cell(row_idx, col_idx)
                value = cell.value
                if value is not None:
                    max_length = max(max_length, len(str(value)))
            ws.column_dimensions[col_letter].width = min(max_length + 2, 28)

        # Row height
        for r in range(header_row, ws.max_row + 1):
            ws.row_dimensions[r].height = 22

    wb = Workbook()

    # -----------------------
    # Time Sheet
    # -----------------------
    ws1 = wb.active
    ws1.title = "Time Sheet"
    time_headers = [
        "Date", "day", "Site ID", "area", "stay place", "Team leader", "Team members",
        "Car (Name/No./Vendor)", "Action details", "work start", "On site duration", "End time",
        "Travel & stay", "travel without stay", "reservation type", "Expense", "Comment", "Manager approval",
    ]
    fill_headers(ws1, "time sheet", profile, month_value, len(time_headers))
    header_row = 6

    for idx, title in enumerate(time_headers, start=1):
        cell = ws1.cell(header_row, idx, title)
        cell.font = header_font
        cell.alignment = center_align
        cell.fill = header_fill
        cell.border = thin_border

    row_idx = header_row + 1
    travel_col = time_headers.index("Travel & stay") + 1
    no_stay_col = time_headers.index("travel without stay") + 1

    for d in days:
        entry = entry_by_date.get(d)

        ws1.cell(row_idx, 1, sheet_date(d))
        ws1.cell(row_idx, 2, sheet_day(d))

        if entry:
            ws1.cell(row_idx, 3, entry.site_id or "")
            ws1.cell(row_idx, 4, entry.area or "")
            ws1.cell(row_idx, 5, entry.stay_place or "")
            ws1.cell(row_idx, 6, entry.team_leader or "")
            ws1.cell(row_idx, 7, entry.team_members_text or "")
            ws1.cell(row_idx, 8, entry.car_details or "")
            ws1.cell(row_idx, 9, entry.action_details or "")
            ws1.cell(row_idx, 10, time_to_sheet(entry.work_start))
            ws1.cell(row_idx, 11, entry.on_site_duration or "")
            ws1.cell(row_idx, 12, time_to_sheet(entry.end_time))

            travel_yes, no_stay_yes = map_travel_flags(entry.travel_stay_type)
            ws1.cell(row_idx, travel_col, travel_yes)
            ws1.cell(row_idx, no_stay_col, no_stay_yes)

            ws1.cell(row_idx, 15, entry.reservation_type or "")
            ws1.cell(row_idx, 16, entry.expense or "")
            ws1.cell(row_idx, 17, entry.comment or "")
            ws1.cell(row_idx, 18, "")

        row_idx += 1

    style_sheet(ws1, header_row, len(time_headers))
    apply_sheet_format(ws1, header_row, len(time_headers))

    # -----------------------
    # Evaluation
    # -----------------------
    ws2 = wb.create_sheet("Evaluation")
    eval_headers = [
        "Date", "day", "Site ID", "Site Type", "Activity", "Team Members QTY", "Tower Type",
        "RRU", "RRU Score", "Fiber", "Fiber Score", "DC", "DC Score",
        "Indoor HU & Survey & PSU", "Indoor Score", "SW ER", "SW ER Score",
        "Power Cabinet", "Power Cabinet Score", "Batteries Addition", "Batteries Addition Score",
        "Score", "Comment",
    ]
    fill_headers(ws2, "evaluation", profile, month_value, len(eval_headers))
    header_row2 = 6

    for idx, title in enumerate(eval_headers, start=1):
        cell = ws2.cell(header_row2, idx, title)
        cell.font = header_font
        cell.alignment = center_align
        cell.fill = header_fill
        cell.border = thin_border

    row_idx = header_row2 + 1
    for d in days:
        entry = entry_by_date.get(d)
        ev = eval_map.get(entry.id) if entry else None

        ws2.cell(row_idx, 1, sheet_date(d))
        ws2.cell(row_idx, 2, sheet_day(d))

        if entry:
            ws2.cell(row_idx, 3, entry.site_id or "")

        if ev:
            ws2.cell(row_idx, 4, getattr(ev, "site_type", "") or "")
            ws2.cell(row_idx, 5, ev.activity or "")
            ws2.cell(row_idx, 6, ev.team_members_qty or "")
            ws2.cell(row_idx, 7, ev.tower_type or "")
            ws2.cell(row_idx, 8, ev.hu_rru_qty if ev.hu_rru_qty is not None else "")
            ws2.cell(row_idx, 9, ev.hu_rru_score if ev.hu_rru_score is not None else "")
            ws2.cell(row_idx, 10, ev.fiber_qty if ev.fiber_qty is not None else "")
            ws2.cell(row_idx, 11, ev.fiber_score if ev.fiber_score is not None else "")
            ws2.cell(row_idx, 12, ev.dc_qty if ev.dc_qty is not None else "")
            ws2.cell(row_idx, 13, ev.dc_score if ev.dc_score is not None else "")
            ws2.cell(row_idx, 14, ev.indoor_hu_survey_psu_qty if ev.indoor_hu_survey_psu_qty is not None else "")
            ws2.cell(row_idx, 15, ev.indoor_hu_survey_psu_score if ev.indoor_hu_survey_psu_score is not None else "")
            ws2.cell(row_idx, 16, ev.sw_er_qty if ev.sw_er_qty is not None else "")
            ws2.cell(row_idx, 17, ev.sw_er_score if ev.sw_er_score is not None else "")
            ws2.cell(row_idx, 18, ev.power_cabinet or "")
            ws2.cell(row_idx, 19, ev.power_cabinet_score if ev.power_cabinet_score is not None else "")
            ws2.cell(row_idx, 20, ev.batteries_addition or "")
            ws2.cell(row_idx, 21, ev.batteries_addition_score if ev.batteries_addition_score is not None else "")
            ws2.cell(row_idx, 22, ev.final_score if ev.final_score is not None else "")
            ws2.cell(row_idx, 23, ev.comment or "")

        row_idx += 1

    style_sheet(ws2, header_row2, len(eval_headers))
    apply_sheet_format(ws2, header_row2, len(eval_headers))

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream

# ---------- Routes ----------
@router.get("/")
def root(request: Request, db: Session = Depends(get_db)):
    profile = get_profile(request, db)
    if not profile_ready(profile):
        return RedirectResponse(url="/setup", status_code=status.HTTP_302_FOUND)
    return RedirectResponse(url="/dashboard", status_code=status.HTTP_302_FOUND)


@router.get("/setup")
def setup_page(request: Request, db: Session = Depends(get_db)):
    profile = get_profile(request, db)
    return templates.TemplateResponse(
        "setup_profile.html",
        {"request": request, "profile": profile, "month_label": current_month_label(), "title": "Setup Profile"},
    )


@router.post("/setup")
def setup_submit(
    request: Request,
    full_name: Annotated[str, Form(...)],
    employee_id: Annotated[str, Form(...)],
    db: Session = Depends(get_db),
):
    emp_id = employee_id.strip()

    # 🔍 شوف هل اليوزر موجود
    profile = db.query(User).filter(User.employee_id == emp_id).first()

    if not profile:
        # 🆕 اعمل يوزر جديد
        profile = User(
            name=full_name.strip(),
            employee_id=emp_id,
            email=f"{emp_id}@app.local",
            password_hash="no_password",
        )
        db.add(profile)
        db.commit()
        db.refresh(profile)
    else:
        # ✏️ حدّث الاسم لو عايز
        profile.name = full_name.strip()
        db.commit()

    # 🍪 خزّن في cookie
    response = RedirectResponse(url="/dashboard", status_code=302)
    response.set_cookie(
        key="profile_id",
        value=str(profile.id),
        max_age=60 * 60 * 24 * 365 * 5,
    )

    return response
    return response

@router.get("/dashboard")
def dashboard(request: Request, db: Session = Depends(get_db)):
    profile = ensure_profile(request, db)
    if isinstance(profile, RedirectResponse):
        return profile
    current_start, current_end = month_bounds(current_month_value())
    recent_entries = db.query(DailyEntry).filter(DailyEntry.user_id == profile.id).order_by(DailyEntry.entry_date.desc()).limit(5).all()
    current_count = (
        db.query(DailyEntry)
        .filter(DailyEntry.user_id == profile.id, DailyEntry.entry_date >= current_start, DailyEntry.entry_date < current_end)
        .count()
    )
    return templates.TemplateResponse(
        "single_dashboard.html",
        {
            "request": request,
            "profile": profile,
            "month_label": current_month_label(),
            "month_value": current_month_value(),
            "recent_entries": recent_entries,
            "current_count": current_count,
            "title": "Home",
        },
    )


@router.get("/submission/new")
def submission_new(request: Request, db: Session = Depends(get_db)):
    profile = ensure_profile(request, db)
    if isinstance(profile, RedirectResponse):
        return profile
    return templates.TemplateResponse(
        "daily_entry_form.html",
        {
            "request": request,
            "profile": profile,
            "month_label": current_month_label(),
            "areas": AREAS,
            "travel_types": TRAVEL_TYPES,
            "reservation_types": RESERVATION_TYPES,
            "action_options": ACTION_OPTIONS,
            "stay_places": STAY_PLACES,
            "today": date.today().isoformat(),
            "form_data": {},
            "error": None,
            "edit_mode": False,
            "entry": None,
            "title": "Time Sheet",
        },
    )


@router.post("/submission/new")
def submission_create(
    request: Request,
    entry_date: Annotated[date, Form(...)],
    site_id: Annotated[str, Form(...)],
    area: Annotated[str, Form(...)],
    team_leader: Annotated[str, Form(...)],
    work_start: Annotated[str, Form(...)],
    end_time: Annotated[str, Form(...)],
    travel_stay_type: Annotated[str, Form(...)],
    reservation_type: Annotated[str, Form(...)],
    expense: Annotated[str, Form(...)],
    team_members_text: Annotated[str, Form(...)],
    car_details: Annotated[str, Form(...)],
    action_details: Annotated[list[str], Form(...)],
    stay_place: Annotated[str | None, Form()] = None,
    comment: Annotated[str | None, Form()] = None,
    db: Session = Depends(get_db),
):
    profile = ensure_profile(request, db)
    if isinstance(profile, RedirectResponse):
        return profile

    start_obj = parse_time_value(work_start)
    end_obj = parse_time_value(end_time)

    form_data = {
        "entry_date": entry_date.isoformat(),
        "site_id": site_id,
        "area": area,
        "stay_place": stay_place or "",
        "team_leader": team_leader,
        "work_start": work_start,
        "end_time": end_time,
        "travel_stay_type": travel_stay_type,
        "reservation_type": reservation_type,
        "expense": expense,
        "team_members_text": team_members_text,
        "car_details": car_details,
        "action_details": action_details,
        "comment": comment or "",
    }

    # Validation
    if not start_obj or not end_obj or end_obj <= start_obj:
        return templates.TemplateResponse(
            "daily_entry_form.html",
            {
                "request": request,
                "profile": profile,
                "month_label": current_month_label(),
                "areas": AREAS,
                "travel_types": TRAVEL_TYPES,
                "reservation_types": RESERVATION_TYPES,
                "action_options": ACTION_OPTIONS,
                "stay_places": STAY_PLACES,
                "today": entry_date.isoformat(),
                "form_data": form_data,
                "error": "End time must be after start time.",
                "edit_mode": False,
                "entry": None,
                "title": "Time Sheet",
            },
            status_code=400,
        )

    if duplicate_entry_exists(db, profile.id, entry_date):
        return templates.TemplateResponse(
            "daily_entry_form.html",
            {
                "request": request,
                "profile": profile,
                "month_label": current_month_label(),
                "areas": AREAS,
                "travel_types": TRAVEL_TYPES,
                "reservation_types": RESERVATION_TYPES,
                "action_options": ACTION_OPTIONS,
                "stay_places": STAY_PLACES,
                "today": entry_date.isoformat(),
                "form_data": form_data,
                "error": "This date is already recorded.",
                "edit_mode": False,
                "entry": None,
                "title": "Time Sheet",
            },
            status_code=400,
        )

    # ✅ إنشاء الريكورد
    entry = DailyEntry(
        user_id=profile.id,
        entry_date=entry_date,
        site_id=site_id.strip(),
        area=area,
        stay_place=stay_place or None,
        team_leader=team_leader.strip(),
        work_start=start_obj,
        end_time=end_obj,
        on_site_duration=calculate_duration_text(start_obj, end_obj),
        travel_stay_type=travel_stay_type,
        reservation_type=reservation_type,
        expense=expense.strip(),
        team_members_text=team_members_text.strip(),
        car_details=car_details.strip(),
        action_details=", ".join(action_details),
        comment=(comment or "").strip() or None,
        submission_status="draft",
        manager_approval="",
    )

    # ✅ أهم سطر (كان الغلط هنا)
    db.add(entry)

    db.commit()
    db.refresh(entry)

    return RedirectResponse(
        url=f"/submission/{entry.id}/evaluation",
        status_code=status.HTTP_302_FOUND
    )


@router.get("/submission/{entry_id}/evaluation")
def evaluation_page(entry_id: int, request: Request, db: Session = Depends(get_db)):
    profile = ensure_profile(request, db)
    if isinstance(profile, RedirectResponse):
        return profile
    entry = db.query(DailyEntry).filter(DailyEntry.id == entry_id, DailyEntry.user_id == profile.id).first()
    if not entry:
        return RedirectResponse(url="/records", status_code=status.HTTP_302_FOUND)
    evaluation = db.query(SiteEvaluation).filter(SiteEvaluation.daily_entry_id == entry.id).first()
    form_data: dict = {}
    calculated = None
    if evaluation:
        form_data = {
            "site_type": evaluation.site_type,
            "activity": (evaluation.activity or "").split(", ") if evaluation.activity else [],
            "team_members_qty": evaluation.team_members_qty,
            "tower_type": evaluation.tower_type,
            "rru_qty": "" if evaluation.hu_rru_qty is None else evaluation.hu_rru_qty,
            "fiber_qty": "" if evaluation.fiber_qty is None else evaluation.fiber_qty,
            "dc_qty": "" if evaluation.dc_qty is None else evaluation.dc_qty,
            "indoor_hu_survey_psu_qty": "" if evaluation.indoor_hu_survey_psu_qty is None else evaluation.indoor_hu_survey_psu_qty,
            "sw_er_qty": "" if evaluation.sw_er_qty is None else evaluation.sw_er_qty,
            "power_cabinet": evaluation.power_cabinet,
            "batteries_addition": evaluation.batteries_addition,
            "comment": evaluation.comment or "",
        }
        calculated = {"final_score": evaluation.final_score or 0}
    return templates.TemplateResponse(
        "site_evaluation_form.html",
        {
            "request": request,
            "profile": profile,
            "entry": entry,
            "evaluation": evaluation,
            "today": entry.entry_date.isoformat(),
            "month_label": current_month_label(),
            "action_options": ACTION_OPTIONS,
            "site_types": SITE_TYPES,
            "yes_no_options": YES_NO,
            "tower_types": TOWER_TYPES,
            "error": None,
            "form_data": form_data,
            "calculated": calculated,
            "title": "Evaluation",
        },
    )


@router.post("/submission/{entry_id}/evaluation")
def evaluation_submit(
    entry_id: int,
    request: Request,
    site_type: Annotated[str, Form(...)],
    activity: Annotated[list[str], Form(...)],
    team_members_qty: Annotated[int, Form(...)],
    tower_type: Annotated[str, Form(...)],
    rru_qty: Annotated[str, Form(...)],
    fiber_qty: Annotated[str, Form(...)],
    dc_qty: Annotated[str, Form(...)],
    indoor_hu_survey_psu_qty: Annotated[str, Form(...)],
    sw_er_qty: Annotated[str, Form(...)],
    power_cabinet: Annotated[str, Form(...)],
    batteries_addition: Annotated[str, Form(...)],
    submit_action: Annotated[str, Form(...)],
    comment: Annotated[str | None, Form()] = None,
    db: Session = Depends(get_db),
):
    profile = ensure_profile(request, db)
    if isinstance(profile, RedirectResponse):
        return profile
    entry = db.query(DailyEntry).filter(DailyEntry.id == entry_id, DailyEntry.user_id == profile.id).first()
    if not entry:
        return RedirectResponse(url="/records", status_code=status.HTTP_302_FOUND)

    payload = {
        "site_type": site_type,
        "team_members_qty": str(team_members_qty),
        "tower_type": tower_type,
        "rru_qty": rru_qty,
        "fiber_qty": fiber_qty,
        "dc_qty": dc_qty,
        "indoor_hu_survey_psu_qty": indoor_hu_survey_psu_qty,
        "sw_er_qty": sw_er_qty,
        "power_cabinet": power_cabinet,
        "batteries_addition": batteries_addition,
    }
    calculated = calc_eval_scores(payload, site_type)
    evaluation = db.query(SiteEvaluation).filter(SiteEvaluation.daily_entry_id == entry.id).first()
    if not evaluation:
        evaluation = SiteEvaluation(
            daily_entry_id=entry.id,
            user_id=profile.id,
            eval_date=entry.entry_date,
            site_id=entry.site_id,
            area=entry.area,
            stay_place=entry.stay_place,
        )
    evaluation.site_type = site_type
    evaluation.activity = ", ".join(activity)
    evaluation.team_members_qty = team_members_qty
    evaluation.tower_type = tower_type
    evaluation.hu_rru_qty = None if rru_qty == "" else to_float(rru_qty)
    evaluation.hu_rru_score = calculated["rru_score"]
    evaluation.er_rru_qty = None
    evaluation.er_rru_score = None
    evaluation.fiber_qty = None if fiber_qty == "" else to_float(fiber_qty)
    evaluation.fiber_score = calculated["fiber_score"]
    evaluation.dc_qty = None if dc_qty == "" else to_float(dc_qty)
    evaluation.dc_score = calculated["dc_score"]
    evaluation.indoor_hu_survey_psu_qty = None if indoor_hu_survey_psu_qty == "" else to_float(indoor_hu_survey_psu_qty)
    evaluation.indoor_hu_survey_psu_score = calculated["indoor_hu_survey_psu_score"]
    evaluation.sw_er_qty = None if sw_er_qty == "" else to_float(sw_er_qty)
    evaluation.sw_er_score = calculated["sw_er_score"]
    evaluation.power_cabinet = power_cabinet
    evaluation.power_cabinet_score = calculated["power_cabinet_score"]
    evaluation.batteries_addition = batteries_addition
    evaluation.batteries_addition_score = calculated["batteries_addition_score"]
    evaluation.violations = None
    evaluation.acceptance = None
    evaluation.performance = None
    evaluation.final_score = calculated["final_score"]
    evaluation.comment = (comment or "").strip() or None
    db.add(evaluation)
    entry.submission_status = "submitted" if submit_action == "submit" else "draft"
    db.add(entry)
    db.commit()
    return RedirectResponse(url="/records", status_code=status.HTTP_302_FOUND)


@router.get("/records")
def records(request: Request, db: Session = Depends(get_db)):
    profile = ensure_profile(request, db)
    if isinstance(profile, RedirectResponse):
        return profile
    entries = db.query(DailyEntry).filter(DailyEntry.user_id == profile.id).order_by(DailyEntry.entry_date.desc()).all()
    grouped = OrderedDict()
    for entry in entries:
        month_name = entry.entry_date.strftime("%B %Y")
        grouped.setdefault(month_name, []).append(entry)
    return templates.TemplateResponse(
        "records_grouped.html",
        {"request": request, "profile": profile, "month_label": current_month_label(), "grouped": grouped, "title": "My Records"},
    )


@router.get("/record/{entry_id}/edit")
def edit_entry_page(entry_id: int, request: Request, db: Session = Depends(get_db)):
    profile = ensure_profile(request, db)
    if isinstance(profile, RedirectResponse):
        return profile
    entry = db.query(DailyEntry).filter(DailyEntry.id == entry_id, DailyEntry.user_id == profile.id).first()
    if not entry:
        return RedirectResponse(url="/records", status_code=status.HTTP_302_FOUND)
    return templates.TemplateResponse(
        "daily_entry_form.html",
        {
            "request": request,
            "profile": profile,
            "month_label": current_month_label(),
            "areas": AREAS,
            "travel_types": TRAVEL_TYPES,
            "reservation_types": RESERVATION_TYPES,
            "action_options": ACTION_OPTIONS,
            "stay_places": STAY_PLACES,
            "today": entry.entry_date.isoformat(),
            "form_data": {
                "entry_date": entry.entry_date.isoformat(),
                "site_id": entry.site_id,
                "area": entry.area,
                "stay_place": entry.stay_place or "",
                "team_leader": entry.team_leader or "",
                "work_start": time_to_input(entry.work_start),
                "end_time": time_to_input(entry.end_time),
                "on_site_duration": entry.on_site_duration or "",
                "travel_stay_type": entry.travel_stay_type or "",
                "reservation_type": entry.reservation_type or "",
                "expense": entry.expense or "",
                "team_members_text": entry.team_members_text or "",
                "car_details": entry.car_details or "",
                "action_details": (entry.action_details or "").split(", ") if entry.action_details else [],
                "comment": entry.comment or "",
            },
            "error": None,
            "edit_mode": True,
            "entry": entry,
            "title": "Edit Time Sheet",
        },
    )


@router.post("/record/{entry_id}/edit")
def edit_entry_submit(
    entry_id: int,
    request: Request,
    entry_date: Annotated[date, Form(...)],
    site_id: Annotated[str, Form(...)],
    area: Annotated[str, Form(...)],
    team_leader: Annotated[str, Form(...)],
    work_start: Annotated[str, Form(...)],
    end_time: Annotated[str, Form(...)],
    travel_stay_type: Annotated[str, Form(...)],
    reservation_type: Annotated[str, Form(...)],
    expense: Annotated[str, Form(...)],
    team_members_text: Annotated[str, Form(...)],
    car_details: Annotated[str, Form(...)],
    action_details: Annotated[list[str], Form(...)],
    stay_place: Annotated[str | None, Form()] = None,
    comment: Annotated[str | None, Form()] = None,
    db: Session = Depends(get_db),
):
    profile = ensure_profile(request, db)
    if isinstance(profile, RedirectResponse):
        return profile
    entry = db.query(DailyEntry).filter(DailyEntry.id == entry_id, DailyEntry.user_id == profile.id).first()
    if not entry:
        return RedirectResponse(url="/records", status_code=status.HTTP_302_FOUND)
    start_obj = parse_time_value(work_start)
    end_obj = parse_time_value(end_time)
    if not start_obj or not end_obj or end_obj <= start_obj:
        return RedirectResponse(url=f"/record/{entry.id}/edit", status_code=status.HTTP_302_FOUND)
    if duplicate_entry_exists(db, profile.id, entry_date, exclude_id=entry.id):
        return RedirectResponse(url=f"/record/{entry.id}/edit", status_code=status.HTTP_302_FOUND)

    entry.entry_date = entry_date
    entry.site_id = site_id.strip()
    entry.area = area
    entry.stay_place = stay_place or None
    entry.team_leader = team_leader.strip()
    entry.work_start = start_obj
    entry.end_time = end_obj
    entry.on_site_duration = calculate_duration_text(start_obj, end_obj)
    entry.travel_stay_type = travel_stay_type
    entry.reservation_type = reservation_type
    entry.expense = expense.strip()
    entry.team_members_text = team_members_text.strip()
    entry.car_details = car_details.strip()
    entry.action_details = ", ".join(action_details)
    entry.comment = (comment or "").strip() or None
    db.add(entry)

    evaluation = db.query(SiteEvaluation).filter(SiteEvaluation.daily_entry_id == entry.id).first()
    if evaluation:
        evaluation.eval_date = entry_date
        evaluation.site_id = entry.site_id
        evaluation.area = entry.area
        evaluation.stay_place = entry.stay_place
        db.add(evaluation)
    db.commit()
    return RedirectResponse(url=f"/submission/{entry.id}/evaluation", status_code=status.HTTP_302_FOUND)

@router.post("/record/{entry_id}/delete")
def delete_entry(entry_id: int, db: Session = Depends(get_db)):
    profile = ensure_profile(request, db)
    if isinstance(profile, RedirectResponse):
        return profile

    entry = db.query(DailyEntry).filter(
        DailyEntry.id == entry_id,
        DailyEntry.user_id == profile.id
    ).first()

    if entry:
        db.query(SiteEvaluation).filter(SiteEvaluation.daily_entry_id == entry.id).delete()
        db.delete(entry)
        db.commit()

    return RedirectResponse(url="/records", status_code=status.HTTP_302_FOUND)


@router.get("/export")
def export_month(month: str | None = None, db: Session = Depends(get_db)):
    profile = ensure_profile(request, db)
    if isinstance(profile, RedirectResponse):
        return profile

    month_value = month or current_month_value()
    start, end = month_bounds(month_value)

    entries = (
        db.query(DailyEntry)
        .filter(
            DailyEntry.user_id == profile.id,
            DailyEntry.entry_date >= start,
            DailyEntry.entry_date < end
        )
        .order_by(DailyEntry.entry_date.asc())
        .all()
    )

    eval_rows = (
        db.query(SiteEvaluation)
        .filter(SiteEvaluation.daily_entry_id.in_([e.id for e in entries]))
        .all()
        if entries else []
    )
    eval_map = {row.daily_entry_id: row for row in eval_rows}

    stream = build_monthly_export(entries, eval_map, profile, month_value)
    filename = f"{profile.name.replace(' ', '_')}_{start.strftime('%Y_%m')}.xlsx"

    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )


@router.get("/settings")
def settings_page(request: Request, db: Session = Depends(get_db)):
    profile = get_profile(request, db)
    return templates.TemplateResponse(
        "setup_profile.html",
        {"request": request, "profile": profile, "month_label": current_month_label(), "is_settings": True, "title": "Settings"},
    )


@router.get("/manifest.json")
def manifest():
    content = {
        "name": "Field Timesheet",
        "short_name": "Timesheet",
        "start_url": "/",
        "display": "standalone",
        "background_color": "#f5f7fb",
        "theme_color": "#1d4ed8",
        "icons": [
            {"src": "/static/icons/icon-192.png", "sizes": "192x192", "type": "image/png"},
            {"src": "/static/icons/icon-512.png", "sizes": "512x512", "type": "image/png"},
        ],
    }
    return Response(content=json.dumps(content), media_type="application/manifest+json")


@router.get("/sw.js")
def sw():
    js = """
const CACHE_NAME = 'field-timesheet-v4';
const STATIC_URLS = [
  '/static/css/style.css'
];

self.addEventListener('install', event => {
  event.waitUntil(
    caches.open(CACHE_NAME).then(cache => cache.addAll(STATIC_URLS))
  );
  self.skipWaiting();
});

self.addEventListener('activate', event => {
  event.waitUntil(
    caches.keys().then(keys =>
      Promise.all(
        keys.map(key => {
          if (key !== CACHE_NAME) {
            return caches.delete(key);
          }
        })
      )
    )
  );
  self.clients.claim();
});

self.addEventListener('fetch', event => {
  if (event.request.method !== 'GET') return;

  const url = new URL(event.request.url);

  // cache static files only
  if (url.pathname.startsWith('/static/')) {
    event.respondWith(
      caches.match(event.request).then(resp => resp || fetch(event.request))
    );
    return;
  }

  // for app pages always go to network first
  event.respondWith(
    fetch(event.request).catch(() => caches.match(event.request))
  );
});
"""
    return Response(content=js, media_type="application/javascript")
